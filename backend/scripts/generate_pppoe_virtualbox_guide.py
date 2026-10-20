"""Build the current guide from its Markdown source, without obsolete API content.

The historical generator name is retained for compatibility. Default output is
install.docx. Convert to legacy .doc with Word/LibreOffice when that format is
required. Render and inspect before distributing generated documents.
"""
import argparse
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'backend/docs/instalasi-ubuntu-vpn.md'


def text_runs(paragraph, text):
    # Preserve literal command text; Markdown inline backticks are presentation.
    for index, part in enumerate(text.split(chr(96))):
        run = paragraph.add_run(part)
        if index % 2:
            run.font.name = 'Consolas'
            run.font.size = Pt(9)


def no_row_split(row):
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement('w:cantSplit'))


def format_table(table):
    properties = table._tbl.tblPr
    for old in properties.findall(qn('w:tblBorders')):
        properties.remove(old)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement('w:' + edge)
        for name, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            border.set(qn('w:' + name), value)
        borders.append(border)
    properties.append(borders)
    for index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_properties = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'E8EEF5' if index == 0 else ('F7F8FA' if index % 2 == 0 else 'FFFFFF'))
            tc_properties.append(shade)
            margins = OxmlElement('w:tcMar')
            for side, size in [('top', '65'), ('bottom', '65'), ('left', '90'), ('right', '90')]:
                margin = OxmlElement('w:' + side)
                margin.set(qn('w:w'), size)
                margin.set(qn('w:type'), 'dxa')
                margins.append(margin)
            tc_properties.append(margins)
            for paragraph in cell.paragraphs:
                # Keep short input tables together and never leave a header alone.
                paragraph.paragraph_format.keep_with_next = index < len(table.rows) - 1 and (len(table.rows) <= 7 or index < 2)
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)


def build(source, destination):
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    normal = document.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in [('Title', 22), ('Heading 1', 14), ('Heading 2', 12)]:
        style = document.styles[name]
        style.font.name = 'Calibri'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        for border in style.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    code_style = document.styles.add_style('Command', 1)
    code_style.font.name = 'Consolas'
    code_style.font.size = Pt(8.3)
    code_style.paragraph_format.line_spacing = 1.0
    code_style.paragraph_format.space_after = Pt(6)
    code_style.paragraph_format.keep_together = True
    document.core_properties.title = 'Instalasi UNZANET di VPS'
    document.core_properties.subject = 'VPS Ubuntu website FreeRADIUS SSTP dan Disconnect'
    document.core_properties.author = 'UNZANET'
    lines = source.read_text(encoding='utf-8-sig').splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.startswith(('~~~', chr(96) * 3)):
            fence = line[:3]
            commands = []
            index += 1
            while index < len(lines) and not lines[index].startswith(fence):
                commands.append(lines[index])
                index += 1
            paragraph = document.add_paragraph('\n'.join(commands), 'Command')
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), 'F2F4F6')
            paragraph._p.get_or_add_pPr().append(shading)
        elif line.startswith('# '):
            document.add_paragraph(line[2:], 'Title')
        elif line.startswith('## '):
            previous = document.paragraphs[-2:]
            if len(previous) == 2 and all(re.match(r'^\d+\. ', paragraph.text) for paragraph in previous):
                previous[0].paragraph_format.keep_with_next = True
            document.add_paragraph(line[3:], 'Heading 1')
        elif line.startswith('### '):
            document.add_paragraph(line[4:], 'Heading 2')
        elif line.startswith('|'):
            rows = []
            while index < len(lines) and lines[index].startswith('|'):
                cells = [cell.strip() for cell in lines[index].strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', cell) for cell in cells):
                    rows.append(cells)
                index += 1
            index -= 1
            table = document.add_table(rows=1, cols=len(rows[0]))
            table.style = 'Table Grid'
            table.autofit = True
            for cell, value in zip(table.rows[0].cells, rows[0]):
                cell.text = value
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            no_row_split(table.rows[0])
            for row_values in rows[1:]:
                row = table.add_row()
                no_row_split(row)
                for cell, value in zip(row.cells, row_values):
                    cell.text = value
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(3)
                        for run in paragraph.runs:
                            run.font.size = Pt(9.5)
            format_table(table)
            document.add_paragraph().paragraph_format.space_after = Pt(0)
        else:
            paragraph = document.add_paragraph()
            if line.endswith(':') or line.startswith('1. Jika '):
                paragraph.paragraph_format.keep_with_next = True
            if line.startswith('- '):
                paragraph.style = 'List Bullet'
                line = line[2:]
            text_runs(paragraph, line)
        index += 1
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    return destination


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, default=SOURCE)
    parser.add_argument('--output', type=Path, default=ROOT / 'install.docx')
    options = parser.parse_args()
    print(build(options.source, options.output))
